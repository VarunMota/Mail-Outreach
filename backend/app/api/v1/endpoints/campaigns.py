from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File, Form
from sqlalchemy.orm import Session
from sqlalchemy import func
from typing import Any, List, Optional
from datetime import datetime, timedelta

from ....db import get_db
from ....models.campaign import Campaign
from ....models.campaign_contact import CampaignContact
from ....models.campaign_attachment import CampaignAttachment
from ....models.contact import Contact
from ....models.email_event import EmailEvent
from ....services.contact_service import ContactService
from ....schemas.analytics import CampaignStats
from ....schemas.campaign import CampaignCreate, CampaignResponse, CampaignListItem

router = APIRouter()


@router.get("/", response_model=List[CampaignListItem])
def list_campaigns(
    skip: int = Query(0, ge=0),
    limit: int = Query(100, ge=1, le=1000),
    db: Session = Depends(get_db)
):
    """
    List all campaigns with their basic stats.
    """
    campaigns = db.query(Campaign).offset(skip).limit(limit).all()
    result = []
    
    for campaign in campaigns:
        # Get stats for each campaign
        stats = db.query(
            func.count(CampaignContact.id).label("total"),
            func.count(CampaignContact.id).filter(CampaignContact.sent_at != None).label("sent"),
            func.count(CampaignContact.id).filter(CampaignContact.opened_at != None).label("opened"),
            func.count(CampaignContact.id).filter(CampaignContact.replied_at != None).label("replied"),
        ).filter(CampaignContact.campaign_id == campaign.id).first()
        
        def safe_div(n, d):
            return round((n / d) * 100, 2) if d > 0 else 0.0
        
        sent_count = stats.sent or 0
        
        result.append({
            "id": campaign.id,
            "name": campaign.name,
            "status": campaign.status,
            "created_at": campaign.created_at,
            "sent_count": sent_count,
            "open_rate": safe_div(stats.opened or 0, sent_count),
            "reply_rate": safe_div(stats.replied or 0, sent_count),
        })
    
    return result


@router.get("/dashboard-summary")
def get_dashboard_summary(db: Session = Depends(get_db)):
    """
    Get aggregated dashboard statistics.
    """
    # Total emails sent
    total_sent = db.query(CampaignContact).filter(CampaignContact.sent_at != None).count()
    
    # Total opened
    total_opened = db.query(CampaignContact).filter(CampaignContact.opened_at != None).count()
    
    # Total clicked
    total_clicked = db.query(CampaignContact).filter(CampaignContact.clicked_at != None).count()
    
    # Total replied
    total_replied = db.query(CampaignContact).filter(CampaignContact.replied_at != None).count()
    
    # Calculate rates
    def safe_div(n, d):
        return round((n / d) * 100, 2) if d > 0 else 0.0
    
    # Recent activity (last 7 days)
    seven_days_ago = datetime.utcnow() - timedelta(days=7)
    
    recent_events = db.query(EmailEvent).filter(
        EmailEvent.timestamp >= seven_days_ago
    ).order_by(EmailEvent.timestamp.desc()).limit(10).all()
    
    activity = []
    for event in recent_events:
        cc = db.query(CampaignContact).get(event.campaign_contact_id)
        if cc and cc.contact:
            activity.append({
                "id": event.id,
                "title": f"Contact {cc.contact.email} - {event.type}",
                "timestamp": event.timestamp.isoformat(),
                "type": event.type
            })
    
    # Daily stats for chart (last 7 days)
    daily_stats = []
    for i in range(6, -1, -1):
        date = datetime.utcnow() - timedelta(days=i)
        day_start = date.replace(hour=0, minute=0, second=0, microsecond=0)
        day_end = day_start + timedelta(days=1)
        
        day_sent = db.query(CampaignContact).filter(
            CampaignContact.sent_at >= day_start,
            CampaignContact.sent_at < day_end
        ).count()
        
        day_opened = db.query(CampaignContact).filter(
            CampaignContact.opened_at >= day_start,
            CampaignContact.opened_at < day_end
        ).count()
        
        day_clicked = db.query(CampaignContact).filter(
            CampaignContact.clicked_at >= day_start,
            CampaignContact.clicked_at < day_end
        ).count()
        
        daily_stats.append({
            "name": date.strftime("%a"),
            "sent": day_sent,
            "opened": day_opened,
            "clicked": day_clicked
        })
    
    return {
        "total_sent": total_sent,
        "total_opened": total_opened,
        "total_clicked": total_clicked,
        "total_replied": total_replied,
        "open_rate": safe_div(total_opened, total_sent),
        "click_rate": safe_div(total_clicked, total_sent),
        "reply_rate": safe_div(total_replied, total_sent),
        "recent_activity": activity,
        "daily_stats": daily_stats
    }


from ....workers.campaign import send_campaign_emails
from ....utils.logging import logger
from ....models.sender_account import SenderAccount
import threading

# Global dictionary to track running campaign threads
running_campaigns = {}

def run_campaign_sync(campaign_id: int):
    """
    Run campaign synchronously in a background thread (no Celery/Redis required).
    """
    try:
        # Import here to avoid circular imports - import the raw function, not the Celery task
        from ....workers.campaign import send_campaign_emails_raw
        
        # Call the task function directly without Celery wrapper
        task_id = f"sync-{campaign_id}"
        result = send_campaign_emails_raw(campaign_id, task_id)
        logger.info(f"[Campaign {campaign_id}] Sync execution completed: {result}")
        
    except Exception as e:
        logger.exception(f"[Campaign {campaign_id}] Error in sync execution: {e}")
    finally:
        # Remove from running campaigns
        if campaign_id in running_campaigns:
            del running_campaigns[campaign_id]

@router.post("/{campaign_id}/start")
def start_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Start or restart a campaign by running it in a background thread.
    Validates sender, template, and contacts before starting.
    Allows restarting completed, stopped, or failed campaigns.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status == "running":
        raise HTTPException(status_code=400, detail="Campaign is already running")
    
    # Allow restart of completed, stopped, or failed campaigns
    # by resetting their contact statuses to pending
    if campaign.status in ["completed", "stopped", "failed"]:
        logger.info(f"[Campaign {campaign_id}] Restarting campaign from {campaign.status} status")
        # Reset all contacts to pending for restart
        db.query(CampaignContact).filter(
            CampaignContact.campaign_id == campaign_id
        ).update({
            "status": "pending",
            "sent_at": None,
            "opened_at": None,
            "clicked_at": None,
            "replied_at": None,
            "bounce_reason": None
        })
        db.commit()
    
    # Validate sender is configured
    if not campaign.sender_id:
        raise HTTPException(
            status_code=400, 
            detail="No sender account configured for this campaign"
        )
    
    sender = db.query(SenderAccount).get(campaign.sender_id)
    if not sender:
        raise HTTPException(
            status_code=400, 
            detail="Configured sender account not found"
        )
    
    if not sender.is_active:
        raise HTTPException(
            status_code=400, 
            detail=f"Sender account '{sender.name}' is inactive"
        )
    
    # Validate campaign has template
    if not campaign.subject_template or not campaign.body_template:
        raise HTTPException(status_code=400, detail="Campaign is missing email template")
    
    # Check if campaign has contacts
    contact_count = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id
    ).count()
    
    if contact_count == 0:
        raise HTTPException(status_code=400, detail="Campaign has no contacts assigned")
    
    # Check for pending contacts
    pending_count = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status == "pending"
    ).count()
    
    if pending_count == 0:
        raise HTTPException(status_code=400, detail="All contacts have already been processed")
    
    # Update status
    campaign.status = "running"
    db.commit()
    
    # Start campaign in background thread (no Celery/Redis needed)
    thread = threading.Thread(target=run_campaign_sync, args=(campaign_id,), daemon=True)
    thread.start()
    running_campaigns[campaign_id] = thread
    
    task_id = f"sync-{campaign_id}"
    campaign.celery_task_id = task_id
    db.commit()
    
    logger.info(f"[Campaign {campaign_id}] Started with sync task, sender={sender.email}")
    
    return {
        "status": "started",
        "campaign_id": campaign_id,
        "task_id": task_id,
        "pending_contacts": pending_count,
        "sender": sender.name
    }


@router.post("/{campaign_id}/stop")
def stop_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Stop a campaign. Worker will exit on next iteration.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status not in ["running", "paused"]:
        raise HTTPException(status_code=400, detail=f"Cannot stop campaign with status '{campaign.status}'")
    
    campaign.status = "stopped"
    db.commit()
    
    logger.info(f"[Campaign {campaign_id}] Stopped by user")
    
    return {"status": "stopped", "campaign_id": campaign_id}


@router.get("/{campaign_id}/stats", response_model=CampaignStats)
def get_campaign_stats(campaign_id: int, db: Session = Depends(get_db)) -> Any:
    """
    Get aggregated statistics for a specific campaign.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")

    # Aggregate counts
    # We can use one query with multiple COUNT(CASE...) or separate queries.
    # Single query is more efficient.
    
    stats = db.query(
        func.count(CampaignContact.id).label("total"),
        func.count(CampaignContact.id).filter(CampaignContact.status == "sent").label("sent"), # Or sent_at is not None? Status is safer.
        # For opened/clicked/replied, we check the timestamps or status?
        # Status might remain 'sent' if we don't update it on open/click (we only log event and update timestamp).
        # We should check the timestamps for accurate engagement counts regardless of current status (which might be 'replied' overriding 'opened')
        func.count(CampaignContact.id).filter(CampaignContact.opened_at != None).label("opened"),
        func.count(CampaignContact.id).filter(CampaignContact.clicked_at != None).label("clicked"),
        func.count(CampaignContact.id).filter(CampaignContact.replied_at != None).label("replied"),
        func.count(CampaignContact.id).filter(CampaignContact.status == "bounced").label("bounced"),
    ).filter(
        CampaignContact.campaign_id == campaign_id
    ).first()

    total = stats.total
    sent = stats.sent
    # If status logic is: pending -> sent -> replied/bounced
    # Then 'sent' status count might decrease as they reply.
    # Better to count sent based on 'sent_at is not None' to get strictly "emails associated with this campaign that were sent"
    
    # Let's refine:
    # Use timestamps/columns that survive status changes.
    real_sent_count = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.sent_at != None
    ).count()
    
    # Recalculate based on real sent count for accuracy
    opened = stats.opened
    clicked = stats.clicked
    replied = stats.replied
    bounced = stats.bounced # Status is reliable for hard bounces that stop flow.

    # Rates
    def safe_div(n, d):
        return round((n / d) * 100, 2) if d > 0 else 0.0

    return CampaignStats(
        total_contacts=total,
        sent=real_sent_count,
        opened=opened,
        clicked=clicked,
        replied=replied,
        bounced=bounced,
        open_rate=safe_div(opened, real_sent_count),
        click_rate=safe_div(clicked, real_sent_count),
        reply_rate=safe_div(replied, real_sent_count),
        bounce_rate=safe_div(bounced, real_sent_count)
    )


@router.get("/{campaign_id}/contacts")
def get_campaign_contacts(
    campaign_id: int, 
    skip: int = 0, 
    limit: int = 20, 
    db: Session = Depends(get_db)
):
    """
    Get list of contacts for a campaign with their status.
    """
    contacts = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id
    ).order_by(CampaignContact.updated_at.desc()).offset(skip).limit(limit).all()
    
    # We can use a schema, but for now let's return a constructed dict or use a generic schema
    # Pydantic schema would be better.
    return [
        {
            "email": c.contact.email, # N+1 problem here, but limit is small (20)
            "status": c.status,
            "sent_at": c.sent_at,
            "opened_at": c.opened_at,
            "clicked_at": c.clicked_at,
            "replied_at": c.replied_at,
            "bounce_reason": c.bounce_reason
        }
        for c in contacts
    ]

@router.post("/{campaign_id}/upload-contacts")
def upload_contacts_to_campaign(
    campaign_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """
    Upload contacts from CSV/Excel file and assign them to a campaign.
    """
    import csv
    import io
    
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    # Validate file type
    allowed_extensions = ['.csv', '.xlsx', '.xls']
    file_ext = '.' + file.filename.split('.')[-1].lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=400, 
            detail=f"Invalid file type. Allowed: {', '.join(allowed_extensions)}"
        )
    
    try:
        # Read file content
        content = file.file.read()
        rows = []
        
        # Parse CSV
        if file_ext == '.csv':
            csv_file = io.StringIO(content.decode('utf-8'))
            reader = csv.DictReader(csv_file)
            rows = list(reader)
        else:
            # Parse Excel files using openpyxl
            from openpyxl import load_workbook
            excel_file = io.BytesIO(content)
            wb = load_workbook(excel_file)
            ws = wb.active
            
            # Get headers from first row
            headers = [cell.value for cell in ws[1]]
            
            # Get data from remaining rows
            for row in ws.iter_rows(min_row=2, values_only=True):
                row_dict = {}
                for i, value in enumerate(row):
                    if i < len(headers):
                        row_dict[headers[i]] = value
                rows.append(row_dict)
        
        # Process contacts
        total_rows = len(rows)
        inserted = 0
        skipped = 0
        invalid = 0
        
        logger.info(f"[Campaign {campaign_id}] Processing {len(rows)} rows from uploaded file")
        logger.info(f"[Campaign {campaign_id}] Column headers found: {list(rows[0].keys()) if rows else 'None'}")
        
        for row in rows:
            # Handle Excel None values and normalize email
            # Support multiple column name variations
            email_val = row.get('email') or row.get('Email') or row.get('EMAIL')
            email = str(email_val).strip() if email_val else ''
            if not email or email.lower() == 'none' or email.lower() == 'nan':
                invalid += 1
                continue
            
            # Normalize other fields (handle None from Excel)
            # Support multiple column name variations
            first_name = row.get('first_name') or row.get('firstName') or row.get('FirstName') or row.get('Firstname') or row.get('FIRSTNAME')
            last_name = row.get('last_name') or row.get('lastName') or row.get('LastName') or row.get('LASTNAME') or row.get('Name')
            company = row.get('company') or row.get('Company') or row.get('COMPANY')
            title = row.get('title') or row.get('Title') or row.get('TITLE')
            linkedin_url = row.get('linkedin_url') or row.get('linkedinUrl') or row.get('LinkedIn') or row.get('Linkedin') or row.get('LINKEDIN')
            
            first_name = str(first_name).strip() if first_name and str(first_name).lower() != 'none' else None
            last_name = str(last_name).strip() if last_name and str(last_name).lower() != 'none' else None
            company = str(company).strip() if company and str(company).lower() != 'none' else None
            title = str(title).strip() if title and str(title).lower() != 'none' else None
            linkedin_url = str(linkedin_url).strip() if linkedin_url and str(linkedin_url).lower() != 'none' else None
            
            # Check if contact already exists
            contact = db.query(Contact).filter(Contact.email == email).first()
            if not contact:
                contact = Contact(
                    email=email,
                    first_name=first_name,
                    last_name=last_name,
                    company=company,
                    title=title,
                    linkedin_url=linkedin_url
                )
                db.add(contact)
                db.commit()
                db.refresh(contact)
                inserted += 1
            else:
                skipped += 1
            
            # Check if already assigned to campaign
            existing = db.query(CampaignContact).filter(
                CampaignContact.campaign_id == campaign_id,
                CampaignContact.contact_id == contact.id
            ).first()
            
            if not existing:
                campaign_contact = CampaignContact(
                    campaign_id=campaign_id,
                    contact_id=contact.id,
                    status='pending'
                )
                db.add(campaign_contact)
        
        db.commit()
        
        return {
            "total_rows": total_rows,
            "inserted": inserted,
            "skipped_duplicates": skipped,
            "invalid_emails": invalid,
            "campaign_id": campaign_id
        }
        
    except Exception as e:
        logger.error(f"Error uploading contacts: {e}")
        raise HTTPException(status_code=500, detail=f"Failed to process file: {str(e)}")

@router.post("/{campaign_id}/pause")
def pause_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Pause a running campaign. Worker will exit cleanly and can be resumed.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status != "running":
        raise HTTPException(status_code=400, detail=f"Cannot pause campaign with status '{campaign.status}'")
    
    campaign.status = "paused"
    db.commit()
    
    logger.info(f"[Campaign {campaign_id}] Paused by user")
    
    return {"status": "paused", "campaign_id": campaign_id}


@router.post("/{campaign_id}/resume")
def resume_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Resume a paused campaign by starting a new background thread.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status != "paused":
        raise HTTPException(status_code=400, detail=f"Cannot resume campaign with status '{campaign.status}'")
    
    # Check for pending contacts
    pending_count = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status == "pending"
    ).count()
    
    if pending_count == 0:
        raise HTTPException(status_code=400, detail="No pending contacts to resume")
    
    # Update status
    campaign.status = "running"
    db.commit()
    
    # Start campaign in background thread (no Celery/Redis needed)
    thread = threading.Thread(target=run_campaign_sync, args=(campaign_id,), daemon=True)
    thread.start()
    running_campaigns[campaign_id] = thread
    
    task_id = f"sync-{campaign_id}"
    campaign.celery_task_id = task_id
    db.commit()
    
    logger.info(f"[Campaign {campaign_id}] Resumed with sync task")
    
    return {
        "status": "resumed",
        "campaign_id": campaign_id,
        "task_id": task_id,
        "pending_contacts": pending_count
    }


@router.get("/{campaign_id}/status")
def get_campaign_status(campaign_id: int, db: Session = Depends(get_db)):
    """
    Get current campaign status with progress information.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    # Get counts
    total = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id
    ).count()
    
    sent = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status == "sent"
    ).count()
    
    pending = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status == "pending"
    ).count()
    
    failed = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status.in_(["failed", "bounced"])
    ).count()
    
    skipped = db.query(CampaignContact).filter(
        CampaignContact.campaign_id == campaign_id,
        CampaignContact.status == "skipped"
    ).count()
    
    percent_complete = round((sent + failed + skipped) / total * 100, 2) if total > 0 else 0
    
    return {
        "campaign_id": campaign_id,
        "status": campaign.status,
        "celery_task_id": campaign.celery_task_id,
        "total_contacts": total,
        "sent_count": sent,
        "pending_count": pending,
        "failed_count": failed,
        "skipped_count": skipped,
        "percent_complete": percent_complete
    }

@router.get("/{campaign_id}")
def get_campaign(campaign_id: int, db: Session = Depends(get_db)):
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    return {
        "id": campaign.id,
        "name": campaign.name,
        "status": campaign.status,
        "created_at": campaign.created_at,
        # "followups": campaign.followups # Add validation model to response
    }

from ....schemas.campaign import CampaignCreate, CampaignResponse
from ....models.followup import CampaignFollowup

@router.post("/", response_model=CampaignResponse)
def create_campaign(
    name: str = Form(...),
    subject_template: str = Form(...),
    body_template: str = Form(...),
    sender_id: int = Form(...),
    daily_limit: int = Form(50),
    schedule: Optional[str] = Form(None),
    attachments: List[UploadFile] = File(default=[]),
    db: Session = Depends(get_db)
):
    """
    Create a new campaign with optional file attachments.
    """
    import os
    import uuid
    from datetime import datetime
    
    # Create upload directory if it doesn't exist
    upload_dir = os.path.join(os.getcwd(), "uploads", "campaigns")
    os.makedirs(upload_dir, exist_ok=True)
    
    # Parse schedule if provided
    scheduled_at = None
    status = "draft"
    if schedule:
        try:
            scheduled_at = datetime.fromisoformat(schedule.replace('Z', '+00:00'))
            status = "scheduled"
        except ValueError:
            pass
    
    # Create Campaign
    campaign = Campaign(
        name=name,
        subject_template=subject_template,
        body_template=body_template,
        daily_limit=daily_limit,
        sender_id=sender_id,
        status=status,
        scheduled_at=scheduled_at
    )
    db.add(campaign)
    db.commit()
    db.refresh(campaign)
    
    # Handle file attachments
    for attachment in attachments:
        if attachment.filename:
            # Generate unique filename
            file_ext = os.path.splitext(attachment.filename)[1]
            unique_filename = f"{uuid.uuid4()}{file_ext}"
            file_path = os.path.join(upload_dir, unique_filename)
            
            # Save file
            with open(file_path, "wb") as f:
                content = attachment.file.read()
                f.write(content)
            
            # Create attachment record
            campaign_attachment = CampaignAttachment(
                campaign_id=campaign.id,
                filename=attachment.filename,
                file_path=file_path,
                file_size=len(content),
                mime_type=attachment.content_type or "application/octet-stream"
            )
            db.add(campaign_attachment)
    
    db.commit()
    db.refresh(campaign)
    
    return campaign


@router.delete("/{campaign_id}")
def delete_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Delete a campaign and all associated data.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    # Delete associated campaign contacts
    db.query(CampaignContact).filter(CampaignContact.campaign_id == campaign_id).delete()
    
    # Delete associated follow-ups
    db.query(CampaignFollowup).filter(CampaignFollowup.campaign_id == campaign_id).delete()
    
    # Delete the campaign
    db.delete(campaign)
    db.commit()
    
    return {"message": "Campaign deleted successfully"}


@router.put("/{campaign_id}", response_model=CampaignResponse)
def update_campaign(
    campaign_id: int,
    campaign_update: CampaignCreate,
    db: Session = Depends(get_db)
):
    """
    Update a campaign. Only allowed for draft or scheduled campaigns.
    """
    from datetime import datetime
    
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    # Only allow editing draft or scheduled campaigns
    if campaign.status not in ["draft", "scheduled"]:
        raise HTTPException(status_code=400, detail="Only draft or scheduled campaigns can be edited")
    
    # Update campaign fields
    campaign.name = campaign_update.name
    campaign.subject_template = campaign_update.subject_template
    campaign.body_template = campaign_update.body_template
    campaign.sender_id = campaign_update.sender_id
    campaign.daily_limit = campaign_update.daily_limit
    
    # Handle schedule update
    if campaign_update.schedule:
        try:
            campaign.scheduled_at = campaign_update.schedule
            campaign.status = "scheduled"
        except ValueError:
            pass
    else:
        campaign.scheduled_at = None
        campaign.status = "draft"
    
    db.commit()
    db.refresh(campaign)
    
    return campaign


@router.get("/{campaign_id}/detail")
def get_campaign_detail(campaign_id: int, db: Session = Depends(get_db)):
    """
    Get detailed campaign information including templates.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    return {
        "id": campaign.id,
        "name": campaign.name,
        "status": campaign.status,
        "subject_template": campaign.subject_template,
        "body_template": campaign.body_template,
        "sender_id": campaign.sender_id,
        "daily_limit": campaign.daily_limit,
        "scheduled_at": campaign.scheduled_at.isoformat() if campaign.scheduled_at else None,
        "created_at": campaign.created_at,
        "updated_at": campaign.updated_at,
    }


@router.post("/{campaign_id}/schedule")
def schedule_campaign(
    campaign_id: int,
    schedule_data: dict,
    db: Session = Depends(get_db)
):
    """
    Schedule a campaign to start at a specific time.
    """
    from datetime import datetime
    
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status not in ["draft", "scheduled"]:
        raise HTTPException(status_code=400, detail="Only draft campaigns can be scheduled")
    
    scheduled_at = schedule_data.get("scheduled_at")
    if not scheduled_at:
        raise HTTPException(status_code=400, detail="scheduled_at is required")
    
    try:
        # Parse ISO format datetime
        campaign.scheduled_at = datetime.fromisoformat(scheduled_at.replace('Z', '+00:00'))
        campaign.status = "scheduled"
        db.commit()
        db.refresh(campaign)
        return {"message": "Campaign scheduled successfully", "scheduled_at": campaign.scheduled_at}
    except ValueError as e:
        raise HTTPException(status_code=400, detail=f"Invalid datetime format: {str(e)}")


@router.post("/{campaign_id}/unschedule")
def unschedule_campaign(campaign_id: int, db: Session = Depends(get_db)):
    """
    Remove schedule from a campaign, reverting it to draft status.
    """
    campaign = db.query(Campaign).get(campaign_id)
    if not campaign:
        raise HTTPException(status_code=404, detail="Campaign not found")
    
    if campaign.status != "scheduled":
        raise HTTPException(status_code=400, detail="Campaign is not scheduled")
    
    campaign.scheduled_at = None
    campaign.status = "draft"
    db.commit()
    db.refresh(campaign)
    return {"message": "Campaign unscheduled successfully", "status": campaign.status}


@router.get("/scheduler/pending")
def get_pending_scheduled_campaigns(db: Session = Depends(get_db)):
    """
    Get all scheduled campaigns that are due to start.
    This endpoint is typically called by a cron job or scheduler service.
    """
    from datetime import datetime
    
    now = datetime.utcnow()
    campaigns = db.query(Campaign).filter(
        Campaign.status == "scheduled",
        Campaign.scheduled_at <= now
    ).all()
    
    return [
        {
            "id": c.id,
            "name": c.name,
            "scheduled_at": c.scheduled_at.isoformat() if c.scheduled_at else None,
        }
        for c in campaigns
    ]
